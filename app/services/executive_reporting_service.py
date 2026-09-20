"""Snapshot-only executive reporting read model and professional file exports."""
from __future__ import annotations

from collections import defaultdict
from datetime import datetime
from io import BytesIO
import re
import unicodedata

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy import and_, or_

from app.models import Product, Representative, RepresentativeBrickAssignment
from app.services.alias_service import AliasService
from app.services.period_service import PeriodService
from app.services.persistent_representative_snapshot_service import (
    PersistentRepresentativeSnapshotService,
)
from app.services.realization_rounding import realization_percent


class ExecutiveReportingService:
    PERIODS = {
        "monthly": "Aylık",
        "quarterly": "3 Aylık",
        "half_year": "6 Aylık · Ocak–Haziran",
        "yearly": "Yıllık",
    }
    SCOPES = {
        "national": "National",
        "region": "Bölge",
        "city": "İl",
        "representative": "Temsilci",
    }
    REGION_LABELS = {
        "101": "İstanbul",
        "201": "Kadıköy",
        "301": "Bursa",
        "401": "İzmir",
        "501": "Ankara",
        "601": "Samsun",
        "602": "Trabzon",
        "701": "Adana",
        "801": "Konya",
        "802": "Antalya",
        "901": "Diyarbakır",
    }

    def __init__(self, *, year=None, month=None, period="monthly", scope="national",
                 scope_value="", product_ids=None):
        active = PeriodService.get_active_period()
        self.year = int(year or active.get("year") or datetime.now().year)
        self.month = max(1, min(12, int(month or active.get("month") or datetime.now().month)))
        self.period = period if period in self.PERIODS else "monthly"
        self.scope = scope if scope in self.SCOPES else "national"
        self.scope_value = str(scope_value or "").strip()
        self.product_ids = {int(value) for value in (product_ids or []) if str(value).isdigit()}

    def months(self):
        if self.period == "monthly":
            return [(self.year, self.month)]
        if self.period == "quarterly":
            ordinal = self.year * 12 + self.month - 1
            return [((ordinal - delta) // 12, (ordinal - delta) % 12 + 1) for delta in (2, 1, 0)]
        if self.period == "half_year":
            return [(self.year, month) for month in range(1, 7)]
        return [(self.year, month) for month in range(1, self.month + 1)]

    @staticmethod
    def _region_code(value):
        match = re.search(r"(?<!\\d)(\\d{3})(?!\\d)", str(value or ""))
        return match.group(1) if match else ""

    @classmethod
    def _region_label(cls, value):
        raw = str(value or "").strip()
        code = cls._region_code(raw)
        descriptive = re.sub(r"^\\s*\\d{3}\\s*", "", raw).strip()
        if descriptive and descriptive != code:
            return descriptive.title()
        return cls.REGION_LABELS.get(code, raw or "Bölge")

    @staticmethod
    def _scope_key(value):
        return AliasService.normalize(value).strip()

    def _assignment_rows(self):
        periods = list(dict.fromkeys(self.months()))
        if not periods:
            return []
        conditions = [
            and_(
                RepresentativeBrickAssignment.year == year,
                RepresentativeBrickAssignment.month == month,
            )
            for year, month in periods
        ]
        return RepresentativeBrickAssignment.query.filter(
            RepresentativeBrickAssignment.active.is_(True),
            or_(*conditions),
        ).all()

    def filter_options(self):
        reps = Representative.query.order_by(
            Representative.region.asc(), Representative.city.asc(), Representative.rep_name.asc()
        ).all()

        region_labels = {}
        for representative in reps:
            raw = str(representative.region or "").strip()
            if not raw:
                continue
            code = self._region_code(raw) or raw
            region_labels.setdefault(code, self._region_label(raw))

        assignments = self._assignment_rows()
        assignment_cities = {
            str(row.city).strip()
            for row in assignments
            if str(row.city or "").strip()
        }
        cities = assignment_cities or {
            str(row.city).strip()
            for row in reps
            if str(row.city or "").strip()
        }

        return {
            "products": Product.query.filter_by(is_active=True).order_by(
                Product.display_order.asc(), Product.product_name.asc()
            ).all(),
            "regions": [
                {"value": value, "label": label}
                for value, label in sorted(
                    region_labels.items(),
                    key=lambda item: (int(item[0]) if str(item[0]).isdigit() else 9999, item[1]),
                )
            ],
            "cities": [
                {"value": city, "label": city}
                for city in sorted(cities, key=self._scope_key)
            ],
            "representatives": reps,
        }

    def _representatives(self):
        rows = Representative.query.order_by(Representative.rep_name.asc()).all()
        if self.scope == "region" and self.scope_value:
            selected_code = self._region_code(self.scope_value) or self.scope_value
            rows = [
                row for row in rows
                if (self._region_code(row.region) or str(row.region or "").strip()) == selected_code
            ]
        elif self.scope == "city" and self.scope_value:
            selected_city = self._scope_key(self.scope_value)
            assignment_rep_ids = {
                int(row.representative_id)
                for row in self._assignment_rows()
                if self._scope_key(row.city) == selected_city
            }
            if assignment_rep_ids:
                rows = [row for row in rows if int(row.id) in assignment_rep_ids]
            else:
                rows = [row for row in rows if self._scope_key(row.city) == selected_city]
        elif self.scope == "representative":
            rows = [
                row for row in rows
                if self.scope_value.isdigit() and int(row.id) == int(self.scope_value)
            ]
        return rows

    @staticmethod
    def _rival_rows(market):
        for row in (market or {}).get("rows") or []:
            product = row.get("product") or {}
            product_id = row.get("product_id") or product.get("id")
            for rival in row.get("rivals") or []:
                yield product_id, rival

    def build(self):
        representatives = self._representatives()
        rep_ids = [row.id for row in representatives]
        product_meta = {row.id: row for row in Product.query.order_by(Product.display_order, Product.product_name).all()}
        buckets = defaultdict(lambda: {
            "target_tl": 0.0, "actual_tl": 0.0, "target_unit": 0.0, "actual_unit": 0.0,
            "market_unit": 0.0, "rivals": defaultdict(float), "months": set(),
        })
        trend = []
        source_weeks = []

        for year, month in self.months():
            snapshots = PersistentRepresentativeSnapshotService.get_active_many(rep_ids, year, month)
            monthly_totals = defaultdict(float)
            for rep_id, payload in snapshots.items():
                monthly = ((payload or {}).get("snapshots") or {}).get("monthly") or {}
                for row in monthly.get("products") or []:
                    product = row.get("product") or {}
                    product_id = row.get("product_id") or product.get("id")
                    if product_id is None:
                        continue
                    product_id = int(product_id)
                    if self.product_ids and product_id not in self.product_ids:
                        continue
                    bucket = buckets[product_id]
                    for key in ("target_tl", "actual_tl", "target_unit", "actual_unit"):
                        value = float(row.get(key, 0) or 0)
                        bucket[key] += value
                        monthly_totals[key] += value
                    bucket["months"].add((year, month))
                market = monthly.get("market_analysis") or {}
                source_week = market.get("source_week") or market.get("latest_week")
                if source_week is not None:
                    source_weeks.append(int(source_week))
                for row in market.get("rows") or []:
                    product = row.get("product") or {}
                    product_id = row.get("product_id") or product.get("id")
                    if product_id is None:
                        continue
                    product_id = int(product_id)
                    if self.product_ids and product_id not in self.product_ids:
                        continue
                    buckets[product_id]["market_unit"] += float(row.get("market_unit", 0) or 0)
                for product_id, rival in self._rival_rows(market):
                    if product_id is None:
                        continue
                    product_id = int(product_id)
                    if self.product_ids and product_id not in self.product_ids:
                        continue
                    name = str(rival.get("name") or rival.get("product_name") or "Rakip").strip()
                    buckets[product_id]["rivals"][name] += float(
                        rival.get("unit", rival.get("value", rival.get("metric_value", 0))) or 0
                    )
            trend.append({
                "label": f"{month:02d}/{year}",
                "actual_tl": round(monthly_totals["actual_tl"], 2),
                "actual_unit": round(monthly_totals["actual_unit"], 2),
            })

        rows = []
        for product_id, values in buckets.items():
            product = product_meta.get(product_id)
            if product is None:
                continue
            market_unit = values["market_unit"]
            company_unit = values["actual_unit"]
            competitor_unit = max(market_unit - company_unit, 0.0) if market_unit else sum(values["rivals"].values())
            rows.append({
                "product_id": product_id,
                "product_name": product.product_name,
                "target_tl": round(values["target_tl"], 2),
                "actual_tl": round(values["actual_tl"], 2),
                "realization_percent": realization_percent(values["actual_tl"], values["target_tl"]) if values["target_tl"] else 0,
                "target_unit": round(values["target_unit"], 2),
                "actual_unit": round(company_unit, 2),
                "market_unit": round(market_unit, 2),
                "competitor_unit": round(competitor_unit, 2),
                "market_share_percent": round(company_unit * 100 / market_unit, 1) if market_unit else None,
                "rivals": [
                    {
                        "name": name,
                        "unit": round(value, 2),
                        "share_percent": round(value * 100 / market_unit, 1) if market_unit else None,
                    }
                    for name, value in sorted(values["rivals"].items(), key=lambda item: item[1], reverse=True)
                ],
            })
        rows.sort(key=lambda row: (product_meta[row["product_id"]].display_order, row["product_name"]))
        totals = {
            key: round(sum(row[key] for row in rows), 2)
            for key in ("target_tl", "actual_tl", "target_unit", "actual_unit", "market_unit", "competitor_unit")
        }
        totals["realization_percent"] = realization_percent(totals["actual_tl"], totals["target_tl"]) if totals["target_tl"] else 0
        totals["market_share_percent"] = round(totals["actual_unit"] * 100 / totals["market_unit"], 1) if totals["market_unit"] else None
        rival_rows = [
            {
                "product_name": row["product_name"], "market_unit": row["market_unit"],
                "company_unit": row["actual_unit"], "company_share_percent": row["market_share_percent"],
                **rival,
            }
            for row in rows for rival in row["rivals"]
        ]
        return {
            "year": self.year, "month": self.month, "period": self.period,
            "period_label": self.PERIODS[self.period], "scope": self.scope,
            "scope_label": self.scope_label(), "rows": rows, "totals": totals,
            "trend": trend, "rival_rows": rival_rows,
            "representative_count": len(representatives),
            "source_week": max(source_weeks) if source_weeks else None,
            "generated_at": datetime.now(),
        }

    def scope_label(self):
        if self.scope == "national":
            return "Türkiye Geneli"
        if self.scope == "region" and self.scope_value:
            return self._region_label(self.scope_value)
        if self.scope == "city" and self.scope_value:
            return self.scope_value
        if self.scope == "representative" and self.scope_value.isdigit():
            row = Representative.query.filter_by(id=int(self.scope_value)).first()
            return row.rep_name if row else "Temsilci"
        return self.SCOPES[self.scope]

    @staticmethod
    def _filename_slug(value):
        text = str(value or "").strip().replace("ı", "i").replace("İ", "I")
        text = unicodedata.normalize("NFKD", text).encode("ascii", "ignore").decode("ascii")
        return re.sub(r"[^a-zA-Z0-9]+", "-", text).strip("-").lower()

    def export_filename(self, report, file_type):
        prefixes = {
            "national": "national-analiz-raporu",
            "region": "bolge-analiz-raporu",
            "city": "il-analiz-raporu",
            "representative": "temsilci-analiz-raporu",
        }
        prefix = prefixes.get(self.scope, "analiz-raporu")
        scope_slug = ""
        if self.scope != "national":
            scope_slug = self._filename_slug(report.get("scope_label"))
        parts = [prefix]
        if scope_slug:
            parts.append(scope_slug)
        parts.extend([str(report["year"]), f'{int(report["month"]):02d}'])
        return "-".join(parts) + f".{file_type}"

    @staticmethod
    def _headers():
        return ["Ürün", "Hedef TL", "Gerçekleşen TL", "Realizasyon %", "Hedef Kutu",
                "Gerçekleşen Kutu", "Toplam Pazar Kutu", "Rakip Kutu", "Pazar Payı %", "Başlıca Rakipler"]

    @staticmethod
    def _excel_title(sheet, title, subtitle, end_column):
        navy, blue, white = "123E70", "0B5CAD", "FFFFFF"
        sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=end_column)
        sheet.cell(1, 1, title)
        sheet.cell(1, 1).font = Font(name="Aptos Display", size=20, bold=True, color=white)
        sheet.cell(1, 1).fill = PatternFill("solid", fgColor=navy)
        sheet.cell(1, 1).alignment = Alignment(vertical="center")
        sheet.row_dimensions[1].height = 34
        sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=end_column)
        sheet.cell(2, 1, subtitle)
        sheet.cell(2, 1).font = Font(name="Aptos", size=10, color=white)
        sheet.cell(2, 1).fill = PatternFill("solid", fgColor=blue)
        sheet.row_dimensions[2].height = 23

    @staticmethod
    def _style_excel_header(cells):
        for cell in cells:
            cell.font = Font(name="Aptos", bold=True, color="FFFFFF", size=10)
            cell.fill = PatternFill("solid", fgColor="0B5CAD")
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    def to_excel(self, report):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Yönetim Özeti"
        subtitle = f'{report["scope_label"]} | {report["period_label"]} | {report["year"]}/{report["month"]:02d}'
        self._excel_title(sheet, "SATIŞ VE PAZAR PERFORMANS RAPORU", subtitle, 10)
        sheet.append([])
        sheet.append(["Kapsam", report["scope_label"], "Dönem", report["period_label"], "Rapor Ayı", f'{report["year"]}/{report["month"]:02d}', "Temsilci", report["representative_count"], "IMS Hafta", report["source_week"] or "-"])
        sheet.append([])
        sheet.append(["Hedef TL", report["totals"]["target_tl"], "Gerçekleşen TL", report["totals"]["actual_tl"], "TL Realizasyon", report["totals"]["realization_percent"] / 100, "Pazar Payı", (report["totals"]["market_share_percent"] or 0) / 100, "Toplam Pazar Kutu", report["totals"]["market_unit"]])
        for column in range(1, 11, 2):
            sheet.cell(6, column).font = Font(name="Aptos", bold=True, color="64748B", size=9)
            sheet.cell(6, column).fill = PatternFill("solid", fgColor="EAF1F8")
            sheet.cell(6, column + 1).font = Font(name="Aptos Display", bold=True, color="123E70", size=13)
            sheet.cell(6, column + 1).fill = PatternFill("solid", fgColor="F5F8FC")
        sheet.cell(6, 2).number_format = '₺#,##0'
        sheet.cell(6, 4).number_format = '₺#,##0'
        sheet.cell(6, 6).number_format = '0%'
        sheet.cell(6, 8).number_format = '0.0%'
        sheet.cell(6, 10).number_format = '#,##0'
        sheet.append([])
        sheet.append(self._headers()[:-1])
        for row in report["rows"]:
            sheet.append([
                row["product_name"], row["target_tl"], row["actual_tl"], row["realization_percent"],
                row["target_unit"], row["actual_unit"], row["market_unit"], row["competitor_unit"],
                row["market_share_percent"],
            ])
        last_row = sheet.max_row
        self._style_excel_header(sheet[8])
        sheet.freeze_panes = "A9"
        sheet.auto_filter.ref = f"A8:I{last_row}"
        thin = Side(style="thin", color="DCE5EF")
        for row in sheet.iter_rows(min_row=9, max_row=last_row, min_col=1, max_col=9):
            for cell in row:
                cell.border = Border(bottom=thin)
                cell.font = Font(name="Aptos", size=10, color="203247")
                cell.alignment = Alignment(vertical="center")
            for cell in row[1:3]: cell.number_format = '₺#,##0'
            row[3].number_format = '0"%"'
            for cell in row[4:8]: cell.number_format = '#,##0'
            row[8].number_format = '0.0"%"'
        widths = [23, 16, 18, 14, 15, 18, 18, 16, 14, 14]
        for index, width in enumerate(widths, 1):
            sheet.column_dimensions[get_column_letter(index)].width = width
        sheet.sheet_view.showGridLines = False
        sheet.print_title_rows = "1:8"
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.print_area = f"A1:J{last_row}"

        trend = workbook.create_sheet("Dönem Trendi")
        self._excel_title(trend, "DÖNEMSEL SATIŞ GELİŞİMİ", subtitle, 4)
        trend.append([]); trend.append(["Dönem", "Gerçekleşen TL", "Gerçekleşen Kutu", "TL Payı"])
        total_trend_tl = sum(float(row["actual_tl"] or 0) for row in report["trend"])
        for item in report["trend"]:
            trend.append([item["label"], item["actual_tl"], item["actual_unit"], (item["actual_tl"] / total_trend_tl) if total_trend_tl else 0])
        self._style_excel_header(trend[4])
        for row in trend.iter_rows(min_row=5, max_row=trend.max_row):
            row[1].number_format = '₺#,##0'; row[2].number_format = '#,##0'; row[3].number_format = '0.0%'
        chart = BarChart(); chart.type = "col"; chart.style = 10; chart.title = "Gerçekleşen TL"; chart.y_axis.title = "TL"; chart.height = 8; chart.width = 16
        chart.add_data(Reference(trend, min_col=2, min_row=4, max_row=trend.max_row), titles_from_data=True)
        chart.set_categories(Reference(trend, min_col=1, min_row=5, max_row=trend.max_row)); trend.add_chart(chart, "F4")
        for index, width in enumerate([18, 20, 20, 14], 1): trend.column_dimensions[get_column_letter(index)].width = width
        trend.freeze_panes = "A5"; trend.sheet_view.showGridLines = False
        trend.page_setup.orientation = "landscape"; trend.page_setup.paperSize = trend.PAPERSIZE_A4
        trend.page_setup.fitToWidth = 1; trend.page_setup.fitToHeight = 1; trend.sheet_properties.pageSetUpPr.fitToPage = True

        rivals = workbook.create_sheet("Rakip Detayı")
        self._excel_title(rivals, "TÜM RAKİPLER VE PAZAR PAYLARI", subtitle, 7)
        rivals.append([]); rivals.append(["Ürün", "Rakip", "Rakip Kutu", "Rakibin Pazar Payı", "Şirket Kutu", "Şirket Pazar Payı", "Toplam Pazar Kutu"])
        for item in report["rival_rows"]:
            rivals.append([item["product_name"], item["name"], item["unit"], (item["share_percent"] or 0) / 100, item["company_unit"], (item["company_share_percent"] or 0) / 100, item["market_unit"]])
        self._style_excel_header(rivals[4]); rivals.freeze_panes = "A5"; rivals.auto_filter.ref = f"A4:G{max(4, rivals.max_row)}"; rivals.sheet_view.showGridLines = False
        for row in rivals.iter_rows(min_row=5, max_row=rivals.max_row):
            row[2].number_format = '#,##0'; row[3].number_format = '0.0%'; row[4].number_format = '#,##0'; row[5].number_format = '0.0%'; row[6].number_format = '#,##0'
        for index, width in enumerate([22, 34, 16, 20, 16, 20, 20], 1): rivals.column_dimensions[get_column_letter(index)].width = width
        rivals.page_setup.orientation = "landscape"; rivals.page_setup.paperSize = rivals.PAPERSIZE_A4
        rivals.page_setup.fitToWidth = 1; rivals.page_setup.fitToHeight = 0; rivals.sheet_properties.pageSetUpPr.fitToPage = True
        rivals.print_title_rows = "1:4"; rivals.print_area = f"A1:G{max(4, rivals.max_row)}"
        output = BytesIO()
        workbook.save(output)
        output.seek(0)
        return output

    def to_pdf(self, report):
        """Create a branded, multi-page management report with full rival detail."""
        from reportlab.lib import colors
        from reportlab.lib.enums import TA_LEFT, TA_RIGHT
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
        from reportlab.lib.units import mm
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
        from reportlab.platypus import KeepTogether, PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

        output = BytesIO()
        font_name, bold_name = "Helvetica", "Helvetica-Bold"
        for regular, bold in (("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"),):
            try:
                pdfmetrics.registerFont(TTFont("ReportSans", regular)); pdfmetrics.registerFont(TTFont("ReportSans-Bold", bold))
                font_name, bold_name = "ReportSans", "ReportSans-Bold"
            except Exception:
                pass
        navy, blue, green, pale, line = colors.HexColor("#123E70"), colors.HexColor("#0B5CAD"), colors.HexColor("#16865B"), colors.HexColor("#F3F7FB"), colors.HexColor("#DCE5EF")
        styles = getSampleStyleSheet()
        title = ParagraphStyle("ReportTitle", parent=styles["Title"], fontName=bold_name, fontSize=19, leading=23, textColor=colors.white, alignment=TA_LEFT)
        subtitle = ParagraphStyle("ReportSubtitle", parent=styles["Normal"], fontName=font_name, fontSize=8.5, leading=12, textColor=colors.HexColor("#DDEBFA"))
        heading = ParagraphStyle("Heading", parent=styles["Heading2"], fontName=bold_name, fontSize=12, leading=15, textColor=navy, spaceAfter=7)
        normal = ParagraphStyle("NormalTR", parent=styles["Normal"], fontName=font_name, fontSize=7.5, leading=10, textColor=colors.HexColor("#25364A"))
        small = ParagraphStyle("SmallTR", parent=normal, fontSize=6.5, leading=8)

        def footer(canvas, doc):
            canvas.saveState(); canvas.setStrokeColor(line); canvas.line(14*mm, 10*mm, 283*mm, 10*mm)
            canvas.setFont(font_name, 6.5); canvas.setFillColor(colors.HexColor("#64748B"))
            canvas.drawString(14*mm, 6*mm, f'IMS Performans Takip Sistemi | {report["scope_label"]}')
            canvas.drawRightString(283*mm, 6*mm, f'Sayfa {doc.page}'); canvas.restoreState()

        doc = SimpleDocTemplate(output, pagesize=landscape(A4), rightMargin=14*mm, leftMargin=14*mm, topMargin=12*mm, bottomMargin=14*mm, title="Satış ve Pazar Performans Raporu")
        story = []
        hero = Table([[Paragraph("SATIŞ VE PAZAR PERFORMANS RAPORU", title), Paragraph(f'{report["scope_label"]}<br/>{report["period_label"]} - {report["year"]}/{report["month"]:02d}', subtitle)]], colWidths=[190*mm, 74*mm])
        hero.setStyle(TableStyle([("BACKGROUND",(0,0),(-1,-1),navy),("BACKGROUND",(1,0),(1,0),blue),("VALIGN",(0,0),(-1,-1),"MIDDLE"),("LEFTPADDING",(0,0),(-1,-1),10),("RIGHTPADDING",(0,0),(-1,-1),10),("TOPPADDING",(0,0),(-1,-1),10),("BOTTOMPADDING",(0,0),(-1,-1),10)])); story += [hero, Spacer(1, 7*mm)]
        total = report["totals"]
        kpis = [["HEDEF TL", "GERÇEKLEŞEN TL", "TL REALİZASYON", "PAZAR PAYI", "TOPLAM PAZAR"], [f'₺{total["target_tl"]:,.0f}', f'₺{total["actual_tl"]:,.0f}', f'%{total["realization_percent"]}', f'%{total["market_share_percent"]:.1f}' if total["market_share_percent"] is not None else "-", f'{total["market_unit"]:,.0f} kutu']]
        kpi_table = Table(kpis, colWidths=[52.8*mm]*5, rowHeights=[7*mm, 13*mm]); kpi_table.setStyle(TableStyle([("BACKGROUND",(0,0),(-1,0),pale),("TEXTCOLOR",(0,0),(-1,0),colors.HexColor("#64748B")),("FONTNAME",(0,0),(-1,0),bold_name),("FONTSIZE",(0,0),(-1,0),6.5),("FONTNAME",(0,1),(-1,1),bold_name),("FONTSIZE",(0,1),(-1,1),13),("TEXTCOLOR",(0,1),(-1,1),navy),("ALIGN",(0,0),(-1,-1),"CENTER"),("VALIGN",(0,0),(-1,-1),"MIDDLE"),("BOX",(0,0),(-1,-1),.5,line),("INNERGRID",(0,0),(-1,-1),.5,line)])); story += [kpi_table, Spacer(1, 7*mm), Paragraph("Ürün performansı", heading)]
        product_data = [["Ürün", "Hedef TL", "Gerçekleşen TL", "Realizasyon", "Hedef Kutu", "Gerçekleşen Kutu", "Toplam Pazar", "Pazar Payı"]]
        for row in report["rows"]:
            product_data.append([Paragraph(row["product_name"], normal), f'₺{row["target_tl"]:,.0f}', f'₺{row["actual_tl"]:,.0f}', f'%{row["realization_percent"]}', f'{row["target_unit"]:,.0f}', f'{row["actual_unit"]:,.0f}', f'{row["market_unit"]:,.0f}', f'%{row["market_share_percent"]:.1f}' if row["market_share_percent"] is not None else "-"])
        product_table = Table(product_data, repeatRows=1, colWidths=[39*mm,35*mm,38*mm,28*mm,31*mm,36*mm,31*mm,27*mm])
        product_table.setStyle(TableStyle([("BACKGROUND",(0,0),(-1,0),blue),("TEXTCOLOR",(0,0),(-1,0),colors.white),("FONTNAME",(0,0),(-1,0),bold_name),("FONTNAME",(0,1),(-1,-1),font_name),("FONTSIZE",(0,0),(-1,-1),7),("ALIGN",(1,1),(-1,-1),"RIGHT"),("ALIGN",(0,0),(-1,0),"CENTER"),("VALIGN",(0,0),(-1,-1),"MIDDLE"),("ROWBACKGROUNDS",(0,1),(-1,-1),[colors.white,pale]),("GRID",(0,0),(-1,-1),.35,line),("TOPPADDING",(0,0),(-1,-1),5),("BOTTOMPADDING",(0,0),(-1,-1),5)])); story += [product_table, PageBreak(), Paragraph("Tüm rakipler ve aylık pazar payları", heading), Paragraph("Pazar payı, seçilen kapsam ve dönemde ilgili ürünün toplam pazar kutusu üzerinden hesaplanır.", normal), Spacer(1, 3*mm)]
        rival_data = [["Ürün", "Rakip", "Rakip Kutu", "Rakibin Pazar Payı", "Şirket Kutu", "Şirket Pazar Payı", "Toplam Pazar"]]
        for item in report["rival_rows"]:
            rival_data.append([Paragraph(item["product_name"], small), Paragraph(item["name"], small), f'{item["unit"]:,.0f}', f'%{item["share_percent"]:.1f}' if item["share_percent"] is not None else "-", f'{item["company_unit"]:,.0f}', f'%{item["company_share_percent"]:.1f}' if item["company_share_percent"] is not None else "-", f'{item["market_unit"]:,.0f}'])
        if len(rival_data) == 1: rival_data.append(["Veri yok", "-", "-", "-", "-", "-", "-"])
        rival_table = Table(rival_data, repeatRows=1, colWidths=[35*mm,74*mm,29*mm,37*mm,29*mm,37*mm,28*mm])
        rival_table.setStyle(TableStyle([("BACKGROUND",(0,0),(-1,0),navy),("TEXTCOLOR",(0,0),(-1,0),colors.white),("FONTNAME",(0,0),(-1,0),bold_name),("FONTNAME",(0,1),(-1,-1),font_name),("FONTSIZE",(0,0),(-1,-1),6.7),("ALIGN",(2,1),(-1,-1),"RIGHT"),("ALIGN",(0,0),(-1,0),"CENTER"),("VALIGN",(0,0),(-1,-1),"MIDDLE"),("ROWBACKGROUNDS",(0,1),(-1,-1),[colors.white,pale]),("GRID",(0,0),(-1,-1),.3,line),("TOPPADDING",(0,0),(-1,-1),4),("BOTTOMPADDING",(0,0),(-1,-1),4)])); story.append(rival_table)
        doc.build(story, onFirstPage=footer, onLaterPages=footer)
        output.seek(0); return output
